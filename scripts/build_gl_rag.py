#!/usr/bin/env python3
"""Clean corpus → data/gl_rag/chunks.jsonl (источник для Node TF-IDF на лендинге)."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "data" / "gl_rag"
SRC = OUT_DIR / "Данные.xlsx"
CHUNKS = OUT_DIR / "chunks.jsonl"

ASZ = re.compile(r"ASZ\s*\d+(?:\s*-\s*\d+)?", re.I)
DEAL = re.compile(r"\b[A-Z]{3}-[A-Z0-9]{3}\b")
SMB = re.compile(r"smb://\S+|\\\\[^\s]+", re.I)
SD = re.compile(r"\bSD\d+\b", re.I)
GENERIC_Q = re.compile(
    r"^(материалы осмотра|непонятны перепланировки|залоговое закл|"
    r"выписки из егрн|установить приоритет|приоритет( по заявке)?|"
    r"отзыв заявки|арбитраж!?|вопрос[ы]? по |знм|зм к сзз|"
    r"отсутствует задание|отсутствует залоговое|в экд нет|"
    r"рм цооп|задание на мониторинг)\b",
    re.I,
)
INTERNAL = re.compile(
    r"не чита(е|ю)т заключение|перевед(ена|ла|ен) на |"
    r"^дубль$|^выставлен\.?$|перевести на ит",
    re.I,
)

TITLE_FIX = {
    "добрый день!": "Актуализация КН — задание на мониторинг не нужно",
    ", вы указали в новой заявке статус осмотра - самоосмотр. при данном": "Статус осмотра самоосмотр / Аспект недоступен в регионе",
    "по заявке asz.... выпущено заключение 03": "Заключение выпущено, сделка не вернулась в СБОФ",
    "в ответ на ваше обращение сообщаем:": "Пересмотр стоимости по ДКП — выставить ВТИ",
}


def ws(s) -> str:
    return "" if s is None else str(s).replace("\xa0", " ").strip()


def mask(s: str) -> str:
    t = ASZ.sub("ASZ....", s)
    t = DEAL.sub("DEAL....", t)
    t = SMB.sub("[ПУТЬ]", t)
    t = SD.sub("SD....", t)
    t = re.sub(r"[ \t]{2,}", " ", t)
    return t.strip()


def first_sentence(s: str) -> str:
    t = re.sub(r"(?i)^добрый день[!.]?\s*", "", s).strip()
    parts = re.split(r"[\n.!?]+", t)
    for p in parts:
        p = p.strip(" \"'")
        if len(p) >= 28:
            return p[:180]
    return t[:180]


def drop_playbook(tema: str, answers: str) -> bool:
    t, a = tema.strip(), answers.strip()
    if len(a) < 20:
        return True
    if t.lower() == "перевести на ит цооп":
        return True
    if INTERNAL.search(a) and len(a) < 80:
        return True
    return False


def fix_title(tema: str, answers: str) -> str:
    key = tema.strip().lower()
    if key in TITLE_FIX:
        return TITLE_FIX[key]
    if len(tema.strip()) < 12 or key.startswith("добрый день"):
        return first_sentence(answers) or tema
    return tema.strip().strip(",. ")


def better_question(q: str, a: str) -> str | None:
    q, a = q.strip(), a.strip()
    if not q or not a:
        return None
    if INTERNAL.search(a) and len(a) < 120:
        return None
    if len(a) < 40:
        return None
    if GENERIC_Q.search(q) or len(q) < 18:
        core = first_sentence(a)
        return core if len(core) >= 28 else None
    return q


def load_src(src: Path):
    wb = load_workbook(src, data_only=True)
    play, gl = [], []
    ws1 = wb["Шаблоны"]
    for i, row in enumerate(ws1.iter_rows(values_only=True)):
        if i == 0:
            continue
        src_name, tema, otv = ws(row[0]), mask(ws(row[1])), mask(ws(row[2]))
        if drop_playbook(tema, otv):
            continue
        tema = fix_title(tema, otv)
        parts = [p.strip() for p in otv.split("\n---\n") if p.strip() and len(p.strip()) >= 20]
        if not parts:
            continue
        play.append({"source": src_name or "Шаблоны", "topic": tema, "answers": parts})
    ws2 = wb["ГЛ"]
    for i, row in enumerate(ws2.iter_rows(values_only=True)):
        if i == 0:
            continue
        q, a = mask(ws(row[0])), mask(ws(row[1]))
        q2 = better_question(q, a)
        if not q2:
            continue
        gl.append({"topic": q2, "answer": a})
    return play, gl


def chunks_from(play, gl):
    out = []
    n = 0
    for p in play:
        for ans in p["answers"]:
            n += 1
            out.append({
                "id": f"tpl-{n:04d}",
                "collection": "template",
                "source": p["source"],
                "topic": p["topic"],
                "text": f"{p['topic']}\n{ans}",
            })
    n = 0
    for g in gl:
        n += 1
        out.append({
            "id": f"case-{n:05d}",
            "collection": "case",
            "source": "ГЛ",
            "topic": g["topic"],
            "text": f"{g['topic']}\n{g['answer']}",
        })
    return out


def norm_key(s: str) -> str:
    t = str(s or "").lower().replace("ё", "е")
    t = re.sub(r"\s+", " ", t).strip()
    return t


STOCK_ANSWER_PREFIXES = (
    "приоритет по заявке выставлен",
    "приоритет по заявку выставлен",
    "приоритет по осмотру выставлен",
    "приоритет по заявке установлен",
    "приоритет обозначен",
)


def question_key(topic: str) -> str:
    t = norm_key(topic)
    for pref in STOCK_ANSWER_PREFIXES:
        if t == pref or t.startswith(pref):
            return pref
    return t


def answer_key(answer: str) -> str:
    """Ключ ответа: первая фраза / шаблон ГЛ в начале ответа."""
    t = norm_key(answer)
    t = re.sub(r"^добрый день[!.]?\s*", "", t)
    for pref in STOCK_ANSWER_PREFIXES:
        idx = t.find(pref)
        if idx >= 0 and idx <= 100:
            return pref
    first_line = t.split("\n", 1)[0].strip()
    first = re.split(r"[.!?]+", first_line, maxsplit=1)[0].strip()
    if len(first) < 24:
        first = first_line[:220] if first_line else t[:220]
    return first[:220]


def split_topic_answer(doc: dict) -> tuple[str, str]:
    topic = str(doc.get("topic") or "").strip()
    text = str(doc.get("text") or "")
    nl = text.find("\n")
    answer = text[nl + 1 :].strip() if nl >= 0 else text.strip()
    if not topic and nl >= 0:
        topic = text[:nl].strip()
    return topic, answer


def dedupe_chunks(docs: list[dict]) -> list[dict]:
    """
    Один чанк на уникальный вопрос И на уникальный ответ (по мягкому ключу).
    Шаблоны и более длинные ответы имеют приоритет.
    """
    def rank(d: dict):
        topic, ans = split_topic_answer(d)
        is_tpl = 0 if d.get("collection") == "template" else 1
        return (is_tpl, -len(ans), -len(topic))

    seen_q: set[str] = set()
    seen_a: set[str] = set()
    kept: list[dict] = []
    for d in sorted(docs, key=rank):
        topic, ans = split_topic_answer(d)
        qk, ak = question_key(topic), answer_key(ans)
        if len(qk) < 8 or len(ak) < 16:
            continue
        if qk in seen_q or ak in seen_a:
            continue
        seen_q.add(qk)
        seen_a.add(ak)
        kept.append(d)

    # стабильные id после схлопывания
    out: list[dict] = []
    ti = ci = 0
    for d in kept:
        nd = dict(d)
        if nd.get("collection") == "template":
            ti += 1
            nd["id"] = f"tpl-{ti:04d}"
        else:
            ci += 1
            nd["id"] = f"case-{ci:05d}"
            nd["collection"] = nd.get("collection") or "case"
        out.append(nd)
    return out


def write_chunks(docs: list[dict], dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", encoding="utf-8") as f:
        for d in docs:
            f.write(json.dumps(d, ensure_ascii=False) + "\n")


def dedupe_chunks_file(src: Path, dest: Path | None = None) -> tuple[int, int]:
    dest = dest or src
    raw = [json.loads(l) for l in src.read_text(encoding="utf-8").splitlines() if l.strip()]
    out = dedupe_chunks(raw)
    write_chunks(out, dest)
    return len(raw), len(out)


def main() -> int:
    # ponytail: `python3 scripts/build_gl_rag.py --dedupe-only` — схлопнуть уже собранный chunks.jsonl
    if len(sys.argv) > 1 and sys.argv[1] in ("--dedupe-only", "-d"):
        target = Path(sys.argv[2]) if len(sys.argv) > 2 else CHUNKS
        if not target.is_file():
            print(f"Нет файла: {target}", file=sys.stderr)
            return 1
        before, after = dedupe_chunks_file(target)
        print(f"dedupe {target}: {before} → {after} (removed {before - after})")
        assert after >= 200
        print("check ok")
        return 0

    src = Path(sys.argv[1]) if len(sys.argv) > 1 else SRC
    if not src.is_file():
        print(f"Нет Excel: {src}", file=sys.stderr)
        print("Положите Данные.xlsx в data/gl_rag/ или передайте путь аргументом.", file=sys.stderr)
        print("Или: python3 scripts/build_gl_rag.py --dedupe-only", file=sys.stderr)
        return 1
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    play, gl = load_src(src)
    docs = dedupe_chunks(chunks_from(play, gl))
    write_chunks(docs, CHUNKS)
    print(f"playbooks {len(play)}  gl {len(gl)}  chunks {len(docs)} → {CHUNKS}")
    assert len(play) >= 40 and len(docs) > 200
    print("check ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
