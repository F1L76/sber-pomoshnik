#!/usr/bin/env python3
"""Пишет марку/модель/год из results.jsonl в xlsx со столбцами VIN / марка / модель / ГоД."""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = ("VIN", "марка", "модель", "ГоД")


def load_results(jsonl_path: Path) -> dict[str, dict]:
    out: dict[str, dict] = {}
    if not jsonl_path.exists():
        return out
    with jsonl_path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            vin = str(row.get("vin") or "").strip().upper()
            if not vin:
                continue
            out[vin] = row
    return out


def _cell(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def write_vin_results(xlsx_path: Path, jsonl_path: Path) -> tuple[int, int]:
    results = load_results(jsonl_path)
    wb = load_workbook(xlsx_path)
    ws = wb.active

    col_vin, col_make, col_model, col_year = 1, 2, 3, 4
    has_headers = False
    for c in range(1, min(ws.max_column + 1, 10)):
        val = str(ws.cell(1, c).value or "").strip().lower()
        if val == "vin":
            col_vin = c
            has_headers = True
        elif "марка" in val or val == "make":
            col_make = c
            has_headers = True
        elif "модель" in val or val == "model":
            col_model = c
            has_headers = True
        elif "год" in val or val == "year":
            col_year = c
            has_headers = True

    if not has_headers:
        for col, title in enumerate(HEADERS, 1):
            ws.cell(1, col, title)

    filled = 0
    for row_i in range(2, ws.max_row + 1):
        vin = str(ws.cell(row_i, col_vin).value or "").strip().upper()
        hit = results.get(vin)
        if not hit or not hit.get("ok"):
            continue
        make = _cell(hit.get("make"))
        model = _cell(hit.get("model"))
        year = _cell(hit.get("year") or hit.get("model_year"))
        if not (make or model or year):
            continue
        if make:
            ws.cell(row_i, col_make, make)
        if model:
            ws.cell(row_i, col_model, model)
        if year:
            ws.cell(row_i, col_year, year if isinstance(year, str) else str(year))
        filled += 1
    wb.save(xlsx_path)
    return filled, len(results)


def _selfcheck() -> None:
    tmp = Path(tempfile.mkdtemp()) / "vin.xlsx"
    jl = tmp.with_suffix(".jsonl")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "VIN"
    ws["B1"] = "марка"
    ws["C1"] = "модель"
    ws["D1"] = "ГоД"
    ws["A2"] = "XTA213100D0147841"
    ws["A3"] = "WBA31AA0X05R89097"
    wb.save(tmp)
    jl.write_text(
        json.dumps(
            {
                "vin": "XTA213100D0147841",
                "ok": True,
                "make": "LADA",
                "model": "2131",
                "year": "2013",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    filled, _ = write_vin_results(tmp, jl)
    wb2 = load_workbook(tmp, data_only=True)
    ws2 = wb2.active
    assert filled == 1, filled
    assert ws2["B2"].value == "LADA"
    assert ws2["C2"].value == "2131"
    assert ws2["D2"].value == "2013"
    assert ws2["B3"].value is None
    print("write-vin-to-xlsx self-check ok")


def main() -> int:
    if len(sys.argv) >= 2 and sys.argv[1] == "--selfcheck":
        _selfcheck()
        return 0
    if len(sys.argv) < 3:
        print("usage: write-vin-to-xlsx.py <xlsx> <results.jsonl>", file=sys.stderr)
        return 2
    filled, total = write_vin_results(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"заполнено строк: {filled} (записей в jsonl: {total})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
